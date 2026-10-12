"""
Profiles the similarity-derived model categories with the INDEPENDENT judge-rubric
performance data (evaluation_results_rubric_v12/all_scores.csv), so each category can
be described by what it is good or bad at -- not just by who it resembles.

The categories come from blind pairwise similarity (analyze_pairwise_clustering.py);
the rubric scores were produced separately by judge_evaluation.py and are never used
in the clustering, so this is a profile of the clusters, not a re-clustering.

Pillars (judge_evaluation.py's seven-criterion rubric; primary zone only, 0-10):
  the SIX full-weight pillars, grouped by what they measure
    Tree reading            Zone identification; Rule extraction (logic rigor)
    Feature analysis        Feature comprehensiveness
    Biological validity     Biological plausibility
    Hypothesis construction Hypothesis integration & specificity; Mechanistic synthesis
  plus Literature grounding, shown separately as a supplementary column because
  compute_final_score() gives it only half weight.

Unit of analysis = the MODEL: each model's score is its mean over the 6 cases, and a
category's value is the mean of its models' means (range = lowest/highest model).
Categories are small (n=3/6/13), so tests are non-parametric (Kruskal-Wallis across the
three categories, Holm-corrected over the six core pillars) and pairwise contrasts are
reported as Cliff's delta with a range-overlap flag; treat p-values as descriptive.

Outputs, under --out-dir (default pairwise_similarity_results_v2/FINAL_3categories_similarity_and_performance/):
    category_profile_summary_table.csv / .xlsx   one row per category: mean (lowest-highest model) per pillar + profile
    model_pillar_means.csv        one row per model: category + mean score per pillar
    category_pillar_summary.csv   category x pillar: mean, SD, min, max, n_models
    category_pillar_tests.csv     Kruskal-Wallis per pillar + pairwise Cliff's delta / Mann-Whitney
    category_profile_summary.txt  the same numbers in readable form
    category_profile_heatmap  .png/.pdf/.svg   category x pillar means (range shown under each mean)
    model_pillar_heatmap      .png/.pdf/.svg   every model x pillar, grouped by category, in dendrogram order

Usage:
    python characterize_categories.py
    python characterize_categories.py --exclude-prompts pRPS6 IL2RA     # 4-case sensitivity check
    python characterize_categories.py --assignments <cluster_assignments.csv> --names "A" "B" "C"
"""

from __future__ import annotations

import argparse
from itertools import combinations
from pathlib import Path

import matplotlib

matplotlib.rcParams["pdf.fonttype"] = 42
matplotlib.rcParams["svg.fonttype"] = "none"  # live text in .svg, editable in Illustrator
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.cm import ScalarMappable
from matplotlib.colors import Normalize
from matplotlib.patches import Rectangle
from scipy.stats import kruskal, mannwhitneyu

from analyze_pairwise_clustering import CLUSTER_PALETTE, model_short_name

SCRIPT_DIR = Path(__file__).resolve().parent
PAIRWISE_DIR = SCRIPT_DIR / "pairwise_similarity_results_v2"
# The final deliverable folder: analyze_pairwise_clustering.py writes the similarity figures and
# cluster_assignments.csv here, and this script adds the performance figures and tables beside them.
FINAL_DIR = PAIRWISE_DIR / "FINAL_3categories_similarity_and_performance"
DEFAULT_ASSIGNMENTS = FINAL_DIR / "cluster_assignments.csv"
DEFAULT_SCORES = SCRIPT_DIR / "evaluation_results_rubric_v12" / "all_scores.csv"
OUT_DIR = FINAL_DIR

# (score column, display label, domain)
CORE_PILLARS = [
    ("Zone Identification score_zone1", "Zone\nidentification", "Tree reading"),
    ("Logic Rigor_zone1", "Rule extraction\n(logic rigor)", "Tree reading"),
    ("Feature Comprehensiveness_zone1", "Feature\ncomprehensiveness", "Feature analysis"),
    ("Biological Plausibility_zone1", "Biological\nplausibility", "Biological validity"),
    ("Specificity_zone1", "Hypothesis integration\n& specificity", "Hypothesis construction"),
    ("Grounded Mechanistic Synthesis_zone1", "Mechanistic\nsynthesis", "Hypothesis construction"),
]
BONUS_PILLAR = ("Literature Grounding_zone1", "Literature grounding\n(bonus, half weight)", "Supplementary")
ALL_PILLARS = CORE_PILLARS + [BONUS_PILLAR]
LABEL_OF = {col: label for col, label, _ in ALL_PILLARS}

# Provisional, evidence-derived names for the default (3-category) assignments; see the
# category_profile_summary.txt numbers they rest on. Other assignment files fall back to
# "Category N" unless --names is given.
DEFAULT_NAMES = {
    1: "Rule-faithful descriptive annotators",
    2: "Integrative, weakly grounded hypothesis generators",
    3: "Mechanistic, literature-grounded integrators",
}

# Qualitative one-line profiles for the default assignments (numbers live in the table columns,
# so this text cannot go stale if the scores are re-averaged over a different case subset).
DEFAULT_PROFILE = {
    1: "Tree reading intact but descriptive: annotates features one by one, with lower feature coverage and "
       "the weakest hypothesis integration and mechanistic synthesis.",
    2: "Builds connected hypotheses (integration close to Category 3) on intact tree reading, but with the "
       "least reliable biology and weakest literature grounding; mechanistic synthesis only moderate.",
    3: "Highest or tied on every pillar, strongest on plausibility, mechanistic synthesis and literature "
       "grounding, though heterogeneous across its models.",
}

CMAP = plt.get_cmap("Blues")
NORM = Normalize(vmin=0, vmax=10)


# ---------------------------------------------------------------------------
# Load + summarise
# ---------------------------------------------------------------------------

def load_model_means(scores_csv: Path, models: list[str], exclude_prompts: list[str]) -> tuple[pd.DataFrame, int]:
    sc = pd.read_csv(scores_csv)
    # all_scores.csv stores the raw model id ("MiniMaxAI/MiniMax-M3"); the pairwise/cluster
    # files use the filesystem slug ("MiniMaxAI__MiniMax-M3"). Normalize before matching.
    sc["model"] = sc["model"].astype(str).str.replace("/", "__", regex=False)
    if exclude_prompts:
        lowered = sc["prompt_file"].astype(str).str.lower()
        sc = sc[~lowered.apply(lambda s: any(x.lower() in s for x in exclude_prompts))]

    missing = sorted(set(models) - set(sc["model"]))
    if missing:
        raise SystemExit(f"ERROR: no evaluation scores for {missing}. Refusing to profile a partial category.")
    sc = sc[sc["model"].isin(models)]
    cases_per_model = sc.groupby("model").size()
    if cases_per_model.nunique() != 1:
        raise SystemExit(f"ERROR: models have unequal case counts:\n{cases_per_model}")

    cols = [c for c, _, _ in ALL_PILLARS]
    means = sc.groupby("model")[cols].mean().reindex(models)
    if means.isna().any().any():
        raise SystemExit("ERROR: NaN pillar means after aggregation (check the score columns).")
    return means, int(cases_per_model.iloc[0])


def category_summary(model_means: pd.DataFrame, cat_of: pd.Series) -> pd.DataFrame:
    rows = []
    for cat, models in cat_of.groupby(cat_of):
        sub = model_means.loc[models.index]
        for col, _, _ in ALL_PILLARS:
            v = sub[col]
            rows.append(dict(category=cat, pillar=col, n_models=len(v), mean=v.mean(),
                             sd=v.std(ddof=1) if len(v) > 1 else np.nan, min=v.min(), max=v.max()))
    return pd.DataFrame(rows)


def holm(p: np.ndarray) -> np.ndarray:
    order = np.argsort(p)
    m = len(p)
    adj = np.empty(m)
    running = 0.0
    for rank, idx in enumerate(order):
        running = max(running, (m - rank) * p[idx])
        adj[idx] = min(1.0, running)
    return adj


def cliffs_delta(a: np.ndarray, b: np.ndarray) -> float:
    gt = (a[:, None] > b[None, :]).sum()
    lt = (a[:, None] < b[None, :]).sum()
    return (gt - lt) / (len(a) * len(b))


def category_tests(model_means: pd.DataFrame, cat_of: pd.Series) -> pd.DataFrame:
    cats = sorted(cat_of.unique())
    groups = {c: cat_of.index[cat_of == c] for c in cats}
    rows = []
    core_p = []
    for col, _, _ in ALL_PILLARS:
        h, p = kruskal(*[model_means.loc[groups[c], col].values for c in cats])
        rows.append(dict(pillar=col, contrast="all categories (Kruskal-Wallis)", statistic=h, p_value=p))
        if col != BONUS_PILLAR[0]:
            core_p.append(p)
        for a, b in combinations(cats, 2):
            xa, xb = model_means.loc[groups[a], col].values, model_means.loc[groups[b], col].values
            u = mannwhitneyu(xa, xb, alternative="two-sided", method="exact") if min(len(xa), len(xb)) else None
            hi, lo = (xa, xb) if xa.mean() >= xb.mean() else (xb, xa)
            rows.append(dict(
                pillar=col, contrast=f"cat {a} vs cat {b}", mean_a=xa.mean(), mean_b=xb.mean(),
                mean_diff_a_minus_b=xa.mean() - xb.mean(), cliffs_delta_a_vs_b=cliffs_delta(xa, xb),
                p_value=u.pvalue if u else np.nan,
                ranges_do_not_overlap=bool(lo.max() < hi.min()),
            ))
    out = pd.DataFrame(rows)
    kw = out["contrast"] == "all categories (Kruskal-Wallis)"
    core = kw & (out["pillar"] != BONUS_PILLAR[0])
    out.loc[core, "p_holm_core_pillars"] = holm(out.loc[core, "p_value"].values)
    return out


def build_summary_table(summary: pd.DataFrame, tests: pd.DataFrame, cat_of: pd.Series, cat_names: dict[int, str],
                        profile: dict[int, str], n_cases: int) -> pd.DataFrame:
    """One row per category ('mean (lowest-highest model)' per pillar) plus a row of between-category p-values."""
    flat = {col: label.replace("\n", " ") for col, label, _ in ALL_PILLARS}
    rows = []
    for cat in sorted(cat_of.unique()):
        models = [model_short_name(m) for m in cat_of.index[cat_of == cat]]
        row = {"Category": cat, "Name": cat_names.get(cat, ""), "n models": len(models), "Models": ", ".join(models)}
        for col, _, _ in ALL_PILLARS:
            r = summary[(summary.category == cat) & (summary.pillar == col)].iloc[0]
            row[flat[col]] = f"{r['mean']:.1f} ({r['min']:.1f}-{r['max']:.1f})"
        row["Profile"] = profile.get(cat, "")
        rows.append(row)
    kw = tests[tests.contrast == "all categories (Kruskal-Wallis)"].set_index("pillar")
    stat_row = {"Category": "Between-category test", "Name": "Kruskal-Wallis, Holm-corrected over the 6 core pillars",
                "n models": len(cat_of), "Models": f"cells = mean of model means over {n_cases} cases (lowest-highest model)"}
    for col, _, _ in ALL_PILLARS:
        r = kw.loc[col]
        stat_row[flat[col]] = (f"p={r['p_holm_core_pillars']:.3f}" if pd.notna(r.get("p_holm_core_pillars"))
                               else f"p={r['p_value']:.3f} (uncorrected, supplementary)")
    stat_row["Profile"] = ""
    rows.append(stat_row)
    return pd.DataFrame(rows)


def save_summary_table(table: pd.DataFrame, out_stem: Path) -> None:
    table.to_csv(out_stem.with_suffix(".csv"), index=False, encoding="utf-8-sig")  # BOM so Excel reads UTF-8
    xlsx = out_stem.with_suffix(".xlsx")
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        table.to_excel(writer, index=False, sheet_name="Category profile")
        ws = writer.sheets["Category profile"]
        from openpyxl.styles import Alignment, Font

        widths = {"Category": 22, "Name": 34, "n models": 9, "Models": 48, "Profile": 60}
        for i, col in enumerate(table.columns, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = widths.get(col, 17)
            ws.cell(row=1, column=i).font = Font(bold=True)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
    print(f"Saved: {out_stem.with_suffix('.csv')} / .xlsx")


# ---------------------------------------------------------------------------
# Plotting
# ---------------------------------------------------------------------------

def _save(fig, out_stem: Path) -> None:
    fig.savefig(out_stem.with_suffix(".png"), dpi=250, bbox_inches="tight")
    fig.savefig(out_stem.with_suffix(".pdf"), bbox_inches="tight")
    fig.savefig(out_stem.with_suffix(".svg"), bbox_inches="tight")
    plt.close(fig)
    print(f"Saved: {out_stem.with_suffix('.png')} / .pdf / .svg")


def _text_color(v: float) -> str:
    return "white" if v >= 6.5 else "#1a1a1a"


def _column_x() -> list[float]:
    """x of each pillar column's left edge: six core columns, a gap, then the bonus column."""
    return list(range(len(CORE_PILLARS))) + [len(CORE_PILLARS) + 0.7]


def _draw_domain_brackets(ax, y_line: float) -> None:
    xs = _column_x()
    spans: dict[str, list[float]] = {}
    for x, (_, _, dom) in zip(xs, ALL_PILLARS):
        spans.setdefault(dom, []).extend([x, x + 1])
    for dom, ends in spans.items():
        lo, hi = min(ends) + 0.06, max(ends) - 0.06
        ax.plot([lo, hi], [y_line, y_line], color="#555555", linewidth=1.0, clip_on=False)
        label = dom if hi - lo > 1.3 else dom.replace(" ", "\n")  # wrap labels over a single column
        ax.text((lo + hi) / 2, y_line - 0.12, label, ha="center", va="bottom", fontsize=9, color="#333333", clip_on=False)


def _draw_column_labels(ax, y: float) -> None:
    for x, (_, label, _) in zip(_column_x(), ALL_PILLARS):
        ax.text(x + 0.5, y, label, ha="center", va="top", fontsize=8.5, color="#222222", clip_on=False)


def _colorbar(fig, ax) -> None:
    cb = fig.colorbar(ScalarMappable(norm=NORM, cmap=CMAP), ax=ax, fraction=0.025, pad=0.05, aspect=18)
    cb.set_label("Mean rubric score (0-10)", fontsize=9)
    cb.solids.set_rasterized(False)
    cb.solids.set_edgecolor("face")


def plot_category_heatmap(
    summary: pd.DataFrame, cat_names: dict[int, str], color_of: dict[int, tuple], out_stem: Path, n_cases: int,
    title_note: str,
) -> None:
    cats = sorted(summary["category"].unique())
    fig, ax = plt.subplots(figsize=(12.5, 1.45 * len(cats) + 3.2))
    for r, cat in enumerate(cats):
        for x, (col, _, _) in zip(_column_x(), ALL_PILLARS):
            row = summary[(summary.category == cat) & (summary.pillar == col)].iloc[0]
            ax.add_patch(Rectangle((x, r), 1, 1, facecolor=CMAP(NORM(row["mean"])), edgecolor="white", linewidth=2.0))
            tc = _text_color(row["mean"])
            ax.text(x + 0.5, r + 0.42, f"{row['mean']:.1f}", ha="center", va="center", fontsize=14, fontweight="bold", color=tc)
            ax.text(x + 0.5, r + 0.76, f"{row['min']:.1f}-{row['max']:.1f}", ha="center", va="center", fontsize=8, color=tc, alpha=0.85)
        n = int(summary[(summary.category == cat)].iloc[0]["n_models"])
        ax.add_patch(Rectangle((-0.42, r + 0.08), 0.16, 0.84, facecolor=color_of[cat], edgecolor="black", linewidth=0.6, clip_on=False))
        ax.text(-0.5, r + 0.5, f"Category {cat} (n={n})\n{cat_names.get(cat, '')}".rstrip(), ha="right", va="center",
                fontsize=9.5, color="#111111", linespacing=1.4, clip_on=False)
    ax.set_xlim(0, _column_x()[-1] + 1)
    ax.set_ylim(len(cats), -0.05)
    ax.axis("off")
    _draw_domain_brackets(ax, -0.2)
    _draw_column_labels(ax, len(cats) + 0.12)
    _colorbar(fig, ax)
    ax.set_title(
        f"Rubric performance profile of the similarity categories\n{title_note}", fontsize=11.5, loc="left", pad=72,
    )
    fig.text(
        0.0, -0.03,
        f"Cell = mean of the category's models' means over {n_cases} cases; small text = lowest-highest model in the category. "
        "Categories are defined by response similarity, not by these scores.",
        fontsize=8, color="#555555",
    )
    _save(fig, out_stem)


def plot_model_heatmap(
    model_means: pd.DataFrame, assignments: pd.DataFrame, cat_names: dict[int, str], color_of: dict[int, tuple],
    out_stem: Path, n_cases: int, title_note: str,
) -> None:
    order = assignments.sort_values("dendrogram_order")["model"].tolist()
    cat_of = assignments.set_index("model")["cluster"]
    n = len(order)
    fig, ax = plt.subplots(figsize=(12.5, 0.36 * n + 3.4))
    prev = None
    for r, model in enumerate(order):
        for x, (col, _, _) in zip(_column_x(), ALL_PILLARS):
            v = model_means.loc[model, col]
            ax.add_patch(Rectangle((x, r), 1, 1, facecolor=CMAP(NORM(v)), edgecolor="white", linewidth=1.2))
            ax.text(x + 0.5, r + 0.5, f"{v:.1f}", ha="center", va="center", fontsize=8.5, color=_text_color(v))
        ax.add_patch(Rectangle((-0.3, r + 0.06), 0.16, 0.88, facecolor=color_of[cat_of[model]], edgecolor="none", clip_on=False))
        ax.text(-0.38, r + 0.5, model_short_name(model), ha="right", va="center", fontsize=8.5, color="#111111", clip_on=False)
        if prev is not None and cat_of[model] != prev:
            ax.plot([-0.3, _column_x()[-1] + 1], [r, r], color="black", linewidth=1.4, clip_on=False, zorder=5)
        prev = cat_of[model]
    ax.set_xlim(0, _column_x()[-1] + 1)
    ax.set_ylim(n, -0.05)
    ax.axis("off")
    _draw_domain_brackets(ax, -0.2)
    _draw_column_labels(ax, n + 0.12)
    _colorbar(fig, ax)
    handles = [Rectangle((0, 0), 1, 1, facecolor=color_of[c], edgecolor="black", linewidth=0.6,
                         label=f"Category {c}: {cat_names.get(c, '')}".rstrip(": ")) for c in sorted(color_of)]
    ax.legend(handles=handles, loc="upper left", bbox_to_anchor=(1.12, 1.0), fontsize=8.5, frameon=False, title="Category (color)")
    ax.set_title(f"Rubric scores per model, grouped by similarity category\n{title_note}", fontsize=11.5, loc="left", pad=62)
    fig.text(0.0, -0.02, f"Each cell = model's mean over {n_cases} cases. Row order matches the similarity dendrogram heatmap.",
             fontsize=8, color="#555555")
    _save(fig, out_stem)


# ---------------------------------------------------------------------------
# Text summary
# ---------------------------------------------------------------------------

def write_summary(path: Path, summary: pd.DataFrame, tests: pd.DataFrame, model_means: pd.DataFrame,
                  cat_of: pd.Series, cat_names: dict[int, str], n_cases: int, title_note: str) -> None:
    short = {c: LABEL_OF[c].replace("\n", " ") for c in LABEL_OF}
    lines = ["CATEGORY PERFORMANCE PROFILE (judge rubric, primary zone, 0-10)", "=" * 78, title_note,
             f"Model score = mean over {n_cases} cases; category value = mean of its models' means.", ""]
    for cat in sorted(cat_of.unique()):
        members = [model_short_name(m) for m in cat_of.index[cat_of == cat]]
        lines.append(f"Category {cat} (n={len(members)}) {cat_names.get(cat, '')}".rstrip())
        lines.append("  " + ", ".join(members))
        sub = summary[summary.category == cat].set_index("pillar")
        for col, _, _ in ALL_PILLARS:
            r = sub.loc[col]
            sd = f"SD {r['sd']:.2f}" if np.isfinite(r["sd"]) else "SD n/a"
            lines.append(f"    {short[col]:40s} mean {r['mean']:.2f}  ({sd}; range {r['min']:.2f}-{r['max']:.2f})")
        lines.append("")
    lines.append("Kruskal-Wallis across categories (unit = model; Holm over the six core pillars):")
    kw = tests[tests.contrast == "all categories (Kruskal-Wallis)"]
    for _, r in kw.iterrows():
        holm_txt = f"Holm p={r['p_holm_core_pillars']:.4f}" if pd.notna(r.get("p_holm_core_pillars")) else "supplementary, not corrected"
        lines.append(f"    {short[r['pillar']]:40s} H={r['statistic']:.2f}  p={r['p_value']:.4f}  ({holm_txt})")
    lines.append("")
    lines.append("Pairwise contrasts (Cliff's delta for the first-named category; +1 = always higher; "
                 "'sep' = model ranges do not overlap; exact Mann-Whitney p is descriptive only):")
    for _, r in tests[tests.contrast.str.startswith("cat ")].iterrows():
        lines.append(f"    {short[r['pillar']]:40s} {r['contrast']:14s} diff={r['mean_diff_a_minus_b']:+.2f}  "
                     f"delta={r['cliffs_delta_a_vs_b']:+.2f}  p={r['p_value']:.4f}  {'sep' if r['ranges_do_not_overlap'] else ''}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Saved: {path}")


# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--assignments", type=str, default=None, help=f"cluster_assignments.csv. Default: {DEFAULT_ASSIGNMENTS}")
    parser.add_argument("--scores", type=str, default=None, help=f"Default: {DEFAULT_SCORES}")
    parser.add_argument("--out-dir", type=str, default=None, help=f"Default: {OUT_DIR}")
    parser.add_argument("--names", nargs="+", default=None,
                        help="Category names in category-id order (default: the provisional names for the default assignments).")
    parser.add_argument("--exclude-prompts", nargs="+", default=[],
                        help="Substrings of prompt_file to leave out of every model's average (e.g. pRPS6 IL2RA).")
    args = parser.parse_args()

    assignments_path = Path(args.assignments) if args.assignments else DEFAULT_ASSIGNMENTS
    out_dir = Path(args.out_dir) if args.out_dir else OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)

    assignments = pd.read_csv(assignments_path)
    models = assignments["model"].tolist()
    cat_of = assignments.set_index("model")["cluster"]
    cats = sorted(cat_of.unique())
    if args.names:
        if len(args.names) != len(cats):
            raise SystemExit(f"ERROR: --names needs {len(cats)} names (one per category {cats}), got {len(args.names)}.")
        cat_names = dict(zip(cats, args.names))
    else:
        cat_names = DEFAULT_NAMES if assignments_path == DEFAULT_ASSIGNMENTS else {}
    color_of = {c: CLUSTER_PALETTE[i % len(CLUSTER_PALETTE)] for i, c in enumerate(cats)}  # same as the cluster map

    model_means, n_cases = load_model_means(Path(args.scores) if args.scores else DEFAULT_SCORES, models, args.exclude_prompts)
    print(f"{len(models)} models x {n_cases} cases each; categories: {cat_of.value_counts().sort_index().to_dict()}")

    summary = category_summary(model_means, cat_of)
    tests = category_tests(model_means, cat_of)
    title_note = f"{len(models)} models, {n_cases} cases" + (f" (excluding prompts: {', '.join(args.exclude_prompts)})" if args.exclude_prompts else "")

    model_means.assign(category=cat_of).to_csv(out_dir / "model_pillar_means.csv")
    summary.to_csv(out_dir / "category_pillar_summary.csv", index=False)
    tests.to_csv(out_dir / "category_pillar_tests.csv", index=False)
    write_summary(out_dir / "category_profile_summary.txt", summary, tests, model_means, cat_of, cat_names, n_cases, title_note)
    profile = DEFAULT_PROFILE if assignments_path == DEFAULT_ASSIGNMENTS else {}
    save_summary_table(build_summary_table(summary, tests, cat_of, cat_names, profile, n_cases), out_dir / "category_profile_summary_table")
    plot_category_heatmap(summary, cat_names, color_of, out_dir / "category_profile_heatmap", n_cases, title_note)
    plot_model_heatmap(model_means, assignments, cat_names, color_of, out_dir / "model_pillar_heatmap", n_cases, title_note)
    print(f"\nDone. Outputs saved under: {out_dir}")


if __name__ == "__main__":
    main()
