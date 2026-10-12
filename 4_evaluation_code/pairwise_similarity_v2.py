"""
Pairwise cross-model similarity/agreement assessment -- v2 rubric.

This is a copy of pairwise_similarity.py with only PAIRWISE_SYSTEM_PROMPT
replaced (the new rubric's OUTPUT FORMAT keys are identical to the original's
CRITERIA list, so no other code changes were needed). Kept as a separate file
so the original pairwise_similarity.py and its cached results are untouched.

This is a similarity metric, not a correctness metric: judge_evaluation.py
already grades each model's answer independently against a biology rubric.
This script instead asks "how much do models X and Y agree/converge on the
same prompt," independent of who is more correct -- two answers can score
high here while both being biologically wrong, and one being more accurate
than the other doesn't by itself lower their similarity.

For every prompt, every pair of models' SELECTED answers (the single answer
select_best_answer.py already picked as most representative for that
model/prompt -- never the raw 3/5 repeats) is sent to a blind LLM judge that
scores six similarity criteria, restricted to how each response treats its
own self-declared PRIMARY zone.

Output layout:
    pairwise_similarity_results_v2/<prompt_stem>/<model_a_slug>__vs__<model_b_slug>.json
        one cached record per (prompt, unordered model pair) -- resumable,
        skipped on re-run unless --overwrite
    pairwise_similarity_results_v2/pairwise_similarity.xlsx
        one tab per model; rows = the prompt cases, columns = every OTHER
        model, cells = Overall Similarity (mean of the 6 criteria) for that
        (this model, other model, prompt) triple, with a red-yellow-green
        conditional-formatting color scale already applied (heatmap-ready).

Usage:
    python pairwise_similarity_v2.py                  # all models x all prompts
    python pairwise_similarity_v2.py --overwrite       # redo cached pairs
    python pairwise_similarity_v2.py --skip-export     # only compute, skip xlsx
    python pairwise_similarity_v2.py --export-only     # only (re)build xlsx from cache
"""

from __future__ import annotations

import argparse
import json
import os
import random
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
TRIAL_RESULTS_DIR = SCRIPT_DIR / "trial_results"
PAIRWISE_DIR = SCRIPT_DIR / "pairwise_similarity_results_v2"
XLSX_OUT = PAIRWISE_DIR / "pairwise_similarity.xlsx"

PAIRWISE_MODEL = "gpt-5.6-sol"

CRITERIA = [
    "Zone & Feature Selection Similarity",
    "Feature Interpretation & Coverage Similarity",
    "Mechanistic Hypothesis Similarity",
    "Functional/Causal-Role Similarity",
    "Supporting Reasoning & Experimental Implication Similarity",
    "Response Scope & Elaboration Similarity",
]

PAIRWISE_SYSTEM_PROMPT = """You are an expert computational biologist evaluating the pairwise similarity between two LLM-generated interpretations of the same decision-tree result.

Response A and Response B are presented in random order. Grade blind and do not infer or mention model identity.

SCOPE

Evaluate only the zone each response identifies as primary, most important, or first-ranked. Ignore secondary and supplementary zones. If a response does not explicitly label a primary zone, use the zone on which it develops its main hypothesis.

GENERAL SCORING RULES

- Evaluate similarity, not correctness or quality. Two responses may be similar even if both are biologically wrong.
- Treat genes, thresholds, branch conditions, and other information supplied by the decision tree as shared input. Their presence alone is not interpretive similarity.
- Award high scores only for comparable interpretive choices: what the features mean, how they are integrated, which causal roles they occupy, and how the evidence supports the conclusion.
- Score each criterion independently. Agreement in zone selection must not automatically raise the biological criteria.
- Compare each claim once and do not reward repetition, answer length, or shared formatting.
- Use integer scores from 0 to 10.

Use these general anchors unless a criterion specifies a stricter ceiling:

- 9-10: Nearly equivalent in the substance evaluated by that criterion.
- 7-8: Strongly similar, with at least one meaningful difference.
- 4-6: Partially similar, with important agreements and important differences.
- 1-3: Only limited, generic, or prompt-inherited overlap.
- 0: No meaningful similarity or directly opposing interpretations.

1. ZONE & FEATURE SELECTION SIMILARITY

Compare the structural decision-tree choice: the primary terminal zone, defining branch conditions, feature directions, and threshold logic. Do not evaluate biological explanations here.

- 9-10: The same terminal zone or genuinely equivalent branch logic, with nearly all defining conditions and directions aligned.
- 7-8: The same zone with a meaningful path discrepancy, or immediately neighboring terminal zones differing by one final split.
- 4-6: Related or partially overlapping branches, but different terminal zones or several important path differences.
- 1-3: Different zones sharing only upstream conditions or a small part of the path.
- 0: Unrelated or opposing branch logic.

Ceilings: Different terminal zones cannot score above 8. If they are not immediately neighboring zones, they cannot score above 6. Shared upstream features cannot override these ceilings.

2. FEATURE INTERPRETATION & COVERAGE SIMILARITY

Compare the context-specific biological interpretations assigned to the focus feature, target, and important condition features, including their functions, directions of effect, and relative emphasis.

- 9-10: Nearly all important features receive the same context-specific interpretations and similar emphasis.
- 7-8: Most important interpretations agree, with one meaningful difference or omission.
- 4-6: Some important interpretations agree, while others differ in function, direction, context, or emphasis.
- 1-3: Overlap is mostly shared feature names, generic background, or one isolated interpretation.
- 0: No substantive agreement or predominantly conflicting interpretations.

Shared feature coverage without shared context-specific interpretation cannot score above 3.

3. MECHANISTIC HYPOTHESIS SIMILARITY

Compare the composed, integrated mechanism rather than individual feature descriptions. Evaluate the central mechanistic conclusion, the major components and intermediate processes, their directed relationships, and how the complete mechanism produces the target outcome.

- 9-10: Essentially the same integrated mechanism, causal architecture, and target explanation.
- 7-8: The same mechanistic backbone and conclusion, with one meaningful difference in an intermediate step, direction, or emphasis.
- 4-6: A related biological theme or conclusion, but substantially different integration or causal architecture.
- 1-3: Only individual feature roles, isolated links, or pathway names overlap; the composed mechanisms differ.
- 0: Fundamentally different, opposing, or absent central mechanisms.

Shared feature functions or pathway vocabulary without a similar integrated mechanism cannot score above 3. Reaching the same target outcome through substantially different mechanisms normally scores 4 or lower.

4. FUNCTIONAL/CAUSAL-ROLE SIMILARITY

Compare how the major components are positioned within the overall hypothesis-for example, as causal drivers, upstream regulators, mediators, permissive contexts, downstream consequences, identity markers, proxies, or correlated readouts. Also compare the strength and certainty of the causal claims.

- 9-10: The major components occupy essentially the same functional and causal positions with similar causal strength.
- 7-8: The overall causal-role structure agrees, with one meaningful difference in role, directness, importance, or certainty.
- 4-6: Some major roles align, but one or more important components are organized differently.
- 1-3: Agreement is limited to generic labels or isolated roles; the overall causal-role structures differ.
- 0: The central role assignments are opposing, unrelated, or undeveloped.

Matching generic labels such as "regulator," "marker," or "driver" without the same role in the integrated hypothesis cannot score above 3.

5. SUPPORTING REASONING & EXPERIMENTAL IMPLICATION SIMILARITY

Compare the full inferential route supporting the primary hypothesis: the biological premises selected, how they bridge the tree pattern to the mechanism, and what causal dependency the proposed experiments or predictions test.

- 9-10: Nearly the same supporting premises, evidence-to-hypothesis reasoning chain, and experimental or predictive logic.
- 7-8: The same main reasoning route and causal test, with one meaningful difference.
- 4-6: Partial overlap in premises or experimental questions, but materially different reasoning bridges, causal contrasts, or predictions.
- 1-3: Only generic biological facts, assay names, or broad experimental goals overlap.
- 0: The reasoning and tests are unrelated, contradictory, or absent.

Shared background facts without a similar reasoning chain cannot score above 3. The same assay counts as similar only if it tests the same mechanistic relationship and predicts a comparable outcome.

6. RESPONSE SCOPE & ELABORATION SIMILARITY
   (Primary-zone discussion only)

Compare the explanatory profile independently of biological agreement: mechanistic resolution, degree of integration, allocation of attention, treatment of alternatives or uncertainty, experimental development, and narrative organization.

- 9-10: Nearly identical depth, synthesis strategy, emphasis, and organization.
- 7-8: Strongly similar explanatory profiles with one noticeable difference.
- 4-6: Broadly comparable level but meaningful differences in depth, integration, emphasis, or organization.
- 1-3: Clearly different explanatory strategies or levels of development.
- 0: Fundamentally different profiles, or one response lacks a developed explanation.

Do not raise this score because the responses select the same zone, discuss the same features, use similar headings, or have similar length. Do not lower it merely because their biological conclusions differ.

FINAL CALIBRATION

For every score above 6, verify that the similarity comes from substantive interpretation rather than shared decision-tree input, terminology, or response structure. Apply all stated ceilings before producing the output.

Overall Similarity is the arithmetic mean of the six scores. The downstream evaluator will calculate it; do not output it.

OUTPUT FORMAT

Output only the six integer scores in a markdown JSON block using exactly these keys:

```json
{
  "Zone & Feature Selection Similarity": 0,
  "Feature Interpretation & Coverage Similarity": 0,
  "Mechanistic Hypothesis Similarity": 0,
  "Functional/Causal-Role Similarity": 0,
  "Supporting Reasoning & Experimental Implication Similarity": 0,
  "Response Scope & Elaboration Similarity": 0
}
```
"""


def load_client(env_var: str = "OPENAI_API_KEY"):
    api_key = os.environ.get(env_var)
    if not api_key:
        raise SystemExit(
            f"ERROR: Environment variable {env_var} is not set.\n"
            'Set it with: [System.Environment]::SetEnvironmentVariable('
            f'"{env_var}", "your_api_key", "User")\n'
            "then restart your terminal."
        )
    from openai import OpenAI

    return OpenAI(api_key=api_key, timeout=300.0)


def model_short_name(model_slug: str) -> str:
    return model_slug.split("__", 1)[-1]


def sheet_name_for(model_slug: str, used: set[str]) -> str:
    """Short, unique Excel sheet name (<=31 chars, no : \\ / ? * [ ])."""
    name = re.sub(r'[:\\/?*\[\]]', "_", model_short_name(model_slug))[:31]
    base = name
    n = 2
    while name in used:
        suffix = f"_{n}"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def load_selected_answers() -> dict[str, dict[str, dict]]:
    """{model_slug: {prompt_stem: {"model": full_model_name, "answer": str}}}
    -- entries with no real answer (selection_method "all_answers_empty", or
    an older null selection) are skipped so they simply don't appear as
    available for pairing on that prompt."""
    result: dict[str, dict[str, dict]] = {}
    for model_dir in sorted(d for d in TRIAL_RESULTS_DIR.iterdir() if d.is_dir()):
        per_prompt = {}
        for sel_path in sorted(model_dir.glob("*_selected_answer.json")):
            data = json.loads(sel_path.read_text(encoding="utf-8"))
            if not data.get("answer"):
                continue
            prompt_stem = Path(data["prompt_file"]).stem
            per_prompt[prompt_stem] = {"model": data["model"], "answer": data["answer"]}
        if per_prompt:
            result[model_dir.name] = per_prompt
    return result


def compare_pair(client, answer_x: str, answer_y: str) -> dict:
    pair = [answer_x, answer_y]
    order = [0, 1]
    random.shuffle(order)  # avoid position bias; scoring is symmetric so no unshuffling needed
    user_content = (
        f"--- RESPONSE A ---\n{pair[order[0]]}\n\n--- RESPONSE B ---\n{pair[order[1]]}"
    )

    response = client.chat.completions.create(
        model=PAIRWISE_MODEL,
        messages=[
            {"role": "system", "content": PAIRWISE_SYSTEM_PROMPT},
            {"role": "user", "content": user_content},
        ],
    )
    raw_text = response.choices[0].message.content or ""
    json_match = re.search(r"```json\s*(.*?)\s*```", raw_text, re.DOTALL)
    clean_json_str = json_match.group(1) if json_match else raw_text.strip().strip("`")
    scores = json.loads(clean_json_str)

    overall = sum((scores.get(c) or 0) for c in CRITERIA) / len(CRITERIA)
    return {"scores": scores, "overall_similarity": overall}


def compute_all(
    client,
    overwrite: bool,
    prompts: list[str] | None = None,
    anchor_models: list[str] | None = None,
    target_models: list[str] | None = None,
    models: list[str] | None = None,
) -> tuple[list[str], list[str]]:
    answers = load_selected_answers()
    model_slugs = sorted(answers.keys())
    if models is not None:
        unknown = set(models) - set(model_slugs)
        if unknown:
            raise SystemExit(f"ERROR: unknown model subfolder(s): {unknown}")
        model_slugs = [m for m in model_slugs if m in set(models)]
    if len(model_slugs) < 2:
        raise SystemExit("ERROR: need at least 2 models with selected answers to compare.")

    all_prompts = sorted({p for m in model_slugs for p in answers[m]})
    if prompts is not None:
        missing = set(prompts) - set(all_prompts)
        if missing:
            raise SystemExit(f"ERROR: unknown prompt stem(s): {missing}")
        all_prompts = [p for p in all_prompts if p in set(prompts)]

    if anchor_models or target_models:
        # Restricted mode: only anchor x target pairs (never anchor x anchor
        # or target x target), so e.g. adding a couple of new models to an
        # already-computed roster doesn't silently blow up into a full new
        # all-pairs recompute.
        if not anchor_models or not target_models:
            raise SystemExit("ERROR: --anchor-models and --target-models must both be given together.")
        unknown = (set(anchor_models) | set(target_models)) - set(model_slugs)
        if unknown:
            raise SystemExit(f"ERROR: unknown model subfolder(s): {unknown}")
        overlap = set(anchor_models) & set(target_models)
        if overlap:
            raise SystemExit(f"ERROR: model(s) listed in both --anchor-models and --target-models: {overlap}")

    done = 0
    skipped = 0
    for prompt_stem in all_prompts:
        available = [m for m in model_slugs if prompt_stem in answers[m]]
        if anchor_models or target_models:
            avail_set = set(available)
            pairs = sorted(
                {
                    tuple(sorted([a, t]))
                    for a in anchor_models
                    for t in target_models
                    if a in avail_set and t in avail_set
                }
            )
        else:
            pairs = [
                tuple(sorted([available[i], available[j]]))
                for i in range(len(available))
                for j in range(i + 1, len(available))
            ]
        for model_a, model_b in pairs:
            out_dir = PAIRWISE_DIR / prompt_stem
            out_dir.mkdir(parents=True, exist_ok=True)
            out_path = out_dir / f"{model_a}__vs__{model_b}.json"

            if out_path.exists() and not overwrite:
                skipped += 1
                continue

            print(f"Comparing [{prompt_stem}] {model_a} vs {model_b} ...", flush=True)
            try:
                result = compare_pair(
                    client,
                    answers[model_a][prompt_stem]["answer"],
                    answers[model_b][prompt_stem]["answer"],
                )
            except Exception as exc:  # noqa: BLE001 - surface the real error, no silent fallback
                print(
                    f"ERROR: comparison failed for [{prompt_stem}] {model_a} vs {model_b}: "
                    f"{type(exc).__name__}: {exc}",
                    file=sys.stderr,
                )
                raise SystemExit(1) from exc

            record = {
                "prompt_file": f"{prompt_stem}.txt",
                "model_a": answers[model_a][prompt_stem]["model"],
                "model_b": answers[model_b][prompt_stem]["model"],
                "model_a_slug": model_a,
                "model_b_slug": model_b,
                "scores": result["scores"],
                "overall_similarity": result["overall_similarity"],
            }
            out_path.write_text(json.dumps(record, indent=2), encoding="utf-8")
            print(f"  -> overall={result['overall_similarity']:.2f}: {out_path}")
            done += 1

    print(f"\nDone. {done} pair(s) compared, {skipped} skipped (cached).")
    return model_slugs, all_prompts


def load_cached_records(all_prompts: list[str]) -> dict[tuple[str, str, str], float]:
    records: dict[tuple[str, str, str], float] = {}
    for prompt_stem in all_prompts:
        d = PAIRWISE_DIR / prompt_stem
        if not d.is_dir():
            continue
        for f in d.glob("*__vs__*.json"):
            data = json.loads(f.read_text(encoding="utf-8"))
            records[(prompt_stem, data["model_a_slug"], data["model_b_slug"])] = data["overall_similarity"]
    return records


def export_workbook(model_slugs: list[str], all_prompts: list[str], out_path: Path | None = None) -> None:
    records = load_cached_records(all_prompts)

    # Restrict the sheet roster to models that actually have a cached pair
    # within this prompt subset, so e.g. a model only ever compared outside
    # this filter (or not compared at all) doesn't show up as an all-blank
    # row/column.
    models_with_data = {m for (_, a, b) in records for m in (a, b)}
    model_slugs = [m for m in model_slugs if m in models_with_data]

    def get_score(prompt_stem: str, m1: str, m2: str) -> float | None:
        a, b = sorted([m1, m2])
        return records.get((prompt_stem, a, b))

    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()

    for focus in model_slugs:
        others = [m for m in model_slugs if m != focus]
        sheet_name = sheet_name_for(focus, used_names)
        ws = wb.create_sheet(title=sheet_name)

        n_cols = 1 + len(others)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
        title_cell = ws.cell(
            row=1, column=1,
            value=f"Model: {model_short_name(focus)} -- Pairwise Overall Similarity (Primary Zone only)",
        )
        title_cell.font = Font(bold=True, size=13)

        header_row = 2
        ws.cell(row=header_row, column=1, value="Prompt File").font = Font(bold=True)
        for c, other in enumerate(others, start=2):
            cell = ws.cell(row=header_row, column=c, value=model_short_name(other))
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")

        first_data_row = header_row + 1
        for r, prompt_stem in enumerate(all_prompts, start=first_data_row):
            ws.cell(row=r, column=1, value=f"{prompt_stem}.txt")
            for c, other in enumerate(others, start=2):
                ws.cell(row=r, column=c, value=get_score(prompt_stem, focus, other))
        last_data_row = first_data_row + len(all_prompts) - 1

        if all_prompts and others:
            rng = f"{get_column_letter(2)}{first_data_row}:{get_column_letter(n_cols)}{last_data_row}"
            ws.conditional_formatting.add(
                rng,
                ColorScaleRule(
                    start_type="num", start_value=0, start_color="F8696B",
                    mid_type="num", mid_value=5, mid_color="FFEB84",
                    end_type="num", end_value=10, end_color="63BE7B",
                ),
            )

        ws.column_dimensions["A"].width = 42
        for c in range(2, n_cols + 1):
            ws.column_dimensions[get_column_letter(c)].width = 18
        ws.freeze_panes = "B3"

    out = out_path if out_path is not None else XLSX_OUT
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Saved: {out}")


def main() -> None:
    global PAIRWISE_DIR, XLSX_OUT

    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--overwrite", action="store_true", help="Redo pairs that already have a cached comparison.")
    parser.add_argument("--skip-export", action="store_true", help="Only compute/cache comparisons, skip writing the xlsx.")
    parser.add_argument("--export-only", action="store_true", help="Skip computing; just (re)build the xlsx from cached comparisons.")
    parser.add_argument(
        "--prompts", nargs="+", default=None,
        help="Subset of prompt stems (no .txt) to compute/export (default: all prompts with any selected answer).",
    )
    parser.add_argument(
        "--models", nargs="+", default=None,
        help="Restrict the full all-pairs roster to these model subfolder names (default: all). "
             "Ignored/combined with --anchor-models/--target-models if those are also given.",
    )
    parser.add_argument(
        "--anchor-models", nargs="+", default=None,
        help="With --target-models: only compute anchor x target pairs (never anchor x anchor or target x target).",
    )
    parser.add_argument(
        "--target-models", nargs="+", default=None,
        help="See --anchor-models.",
    )
    parser.add_argument("--out", type=str, default=None, help=f"Output .xlsx path (default: {XLSX_OUT}).")
    parser.add_argument(
        "--results-dir", type=str, default=None,
        help=f"Directory for cached per-pair JSON + default xlsx location (default: {PAIRWISE_DIR}). "
             "Use a fresh folder to test a rubric revision without touching the existing cache.",
    )
    parser.add_argument(
        "--api-key-env", type=str, default="OPENAI_API_KEY",
        help="Name of the environment variable holding the OpenAI API key to use "
             "(default: OPENAI_API_KEY). Use this to switch between multiple keys, "
             "e.g. --api-key-env OPENAI_API_KEY_2.",
    )
    args = parser.parse_args()

    if not TRIAL_RESULTS_DIR.is_dir():
        raise SystemExit(f"ERROR: trial_results folder not found: {TRIAL_RESULTS_DIR}")

    if args.results_dir:
        PAIRWISE_DIR = Path(args.results_dir)
        XLSX_OUT = PAIRWISE_DIR / "pairwise_similarity.xlsx"

    out_path = Path(args.out) if args.out else None

    if args.export_only:
        answers = load_selected_answers()
        model_slugs = sorted(answers.keys())
        if args.models is not None:
            model_slugs = [m for m in model_slugs if m in set(args.models)]
        all_prompts = sorted({p for m in model_slugs for p in answers[m]})
        if args.prompts is not None:
            all_prompts = [p for p in all_prompts if p in set(args.prompts)]
        export_workbook(model_slugs, all_prompts, out_path)
        return

    client = load_client(args.api_key_env)
    model_slugs, all_prompts = compute_all(
        client, args.overwrite,
        prompts=args.prompts,
        anchor_models=args.anchor_models,
        target_models=args.target_models,
        models=args.models,
    )

    if not args.skip_export:
        export_workbook(model_slugs, all_prompts, out_path)


if __name__ == "__main__":
    main()
